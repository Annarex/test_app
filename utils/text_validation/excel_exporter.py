"""
Модуль для экспорта данных в Excel с подсветкой.
Содержит функции для создания Excel файлов с rich text форматированием.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from typing import List, Optional
from .text_highlighter import create_highlighted_text
from .code_processor import normalize_classification_code
from .error_finder import ErrorInfo


def create_error_report(
    errors: List[ErrorInfo],
    klass_codes: List[str],
    reference_codes: List[str],
    reference_names: List[str] = None,
    headers: List[str] = None
) -> Workbook:
    """
    Создает Excel файл с отчетом об ошибках.
    
    Args:
        errors: Список объектов ErrorInfo с информацией об ошибках
        klass_codes: Коды из классификации
        reference_codes: Коды из справочника
        reference_names: Наименования из справочника (для вывода эталона при ошибках в коде)
        headers: Заголовки столбцов (по умолчанию стандартные)
    
    Returns:
        Workbook объект
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Ошибочные"

    # Заголовки
    if headers is None:
        headers = [
            "Классифицируемый текст",
            "Эталон из справочника",
            "Distance",
            "Код классификации",
            "Код справочника"
        ]
    
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(b=True)

    row = 2
    for error in errors:
        # A: ОРИГИНАЛ с подсветкой отличий
        highlighted = create_highlighted_text(
            error.original_text,
            error.diff_indices,
            error.corrections
        )
        ws.cell(row=row, column=1, value=highlighted)

        # B: эталон (как есть)
        # Если есть ошибка в коде, но нет эталона в наименовании, используем наименование из справочника по индексу кода
        ref_text = error.reference_text
        if error.code_error and not ref_text and reference_names:
            if error.code_error.code_ref_idx is not None and error.code_error.code_ref_idx < len(reference_names):
                ref_text = reference_names[error.code_error.code_ref_idx]
        ws.cell(row=row, column=2, value=ref_text)

        # C: расстояние Левенштейна по ОРИГИНАЛАМ
        ws.cell(row=row, column=3, value=error.distance)

        # D: код из классификации с подсветкой ошибок (если есть ошибки)
        klass_val = klass_codes[error.original_index] if error.original_index < len(klass_codes) else ""
        if error.code_error:
            # Нормализуем код для подсветки (убираем первые 3 символа, если длина = 20)
            normalized_code = normalize_classification_code(klass_val)
            code_highlighted = create_highlighted_text(
                normalized_code,
                error.code_error.diff_indices,
                error.code_error.corrections
            )
            ws.cell(row=row, column=4, value=code_highlighted)
        else:
            # Если нет ошибок в коде, показываем обычный код
            ws.cell(row=row, column=4, value=klass_val)

        # E: код из справочника
        # Используем индекс из code_error, если есть ошибка в коде, иначе используем reference_index
        code_idx_to_use = None
        if error.code_error:
            code_idx_to_use = error.code_error.code_ref_idx
        else:
            code_idx_to_use = error.reference_index
        
        reference_val = reference_codes[code_idx_to_use] if code_idx_to_use is not None and code_idx_to_use < len(reference_codes) else ""
        ws.cell(row=row, column=5, value=reference_val)

        row += 1

    # Автоширина и форматирование
    _apply_column_formatting(ws)
    
    return wb


def _apply_column_formatting(ws):
    """Применяет автоширину и перенос текста для всех столбцов."""
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        
        # Проходим по столбцу один раз для вычисления max_len и установки форматирования
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
            # Устанавливаем перенос текста для всех ячеек
            cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
        
        # Устанавливаем ширину столбца (максимум 50)
        ws.column_dimensions[col_letter].width = min(max_len + 2, 50)


def save_workbook(wb: Workbook, filename: str):
    """
    Сохраняет Workbook в файл.
    
    Args:
        wb: Workbook объект
        filename: Имя файла для сохранения
    """
    wb.save(filename)

